"""Value averaging month by month as a live Excel workbook.

Only the NSE inputs (date, Nifty close, P/E) and the settings are values;
target, gap, investment, units, value, the same-budget version, the fixed SIP
comparison, the yearly roll-up and every XIRR are Excel formulas, so changing
a setting (SIP, target growth, cap, liquid rate) re-runs the whole thing.

    python research/scripts/render_value_averaging_xlsx.py

Needs openpyxl and research/data/smart_sip/value_averaging_ledger.csv (from
value_averaging_ledger.py). Open it in Excel (or recalculate) to see values.
Out: docs/reports/smart-sip/value_averaging_monthly.xlsx
"""
import csv
import os
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LEDGER = os.path.join(ROOT, "data", "smart_sip", "value_averaging_ledger.csv")
NSE_CSV = os.path.join(ROOT, "data", "nifty50_nse_daily_1999.csv")
OUT = os.path.join(os.path.dirname(ROOT), "docs", "reports", "smart-sip", "value_averaging_monthly.xlsx")

ARIAL = "Arial"
BLUE = Font(name=ARIAL, color="0000FF")
BLACK = Font(name=ARIAL)
BOLD = Font(name=ARIAL, bold=True)
HEAD = Font(name=ARIAL, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
GROUP_FILL = {"va": "DCE6F1", "fair": "E2EFDA", "fixed": "FCE4D6"}
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))
RUPEE = '"₹"#,##0;("₹"#,##0);"-"'
PCT = "0.00%"


def last_close():
    with open(NSE_CSV) as fh:
        rows = list(csv.DictReader(fh))
    return date.fromisoformat(rows[-1]["date"]), float(rows[-1]["close"])


def settings_sheet(wb, val_date, val_close):
    ws = wb.active
    ws.title = "Settings"
    ws["A1"] = "Value averaging on Nifty 50 - settings"
    ws["A1"].font = Font(name=ARIAL, bold=True, size=14)
    ws["A2"] = "Yellow cells are inputs; change them and every sheet recalculates."
    ws["A2"].font = Font(name=ARIAL, italic=True, color="666666")
    items = [
        (4, "Monthly SIP / budget (₹)", 10000, RUPEE, "User's base SIP"),
        (5, "Target growth per month", 0.01, PCT, "Target path grows 1% a month (~12.7% a year)"),
        (6, "Max investment per month (₹)", 30000, RUPEE, "3x the SIP"),
        (7, "Liquid fund rate (per year)", 0.06, PCT, "Idle cash in the same-budget version"),
        (8, "Liquid fund rate (per month)", "=(1+B7)^(1/12)-1", "0.0000%", "Formula from B7"),
        (10, "Valuation date", val_date, "dd-mmm-yyyy", "Last NSE session in the data"),
        (11, "Nifty 50 close on valuation date", val_close, "#,##0.00", "Source: NSE historicalOR/indicesHistory"),
    ]
    for r, label, value, fmt, note in items:
        ws.cell(r, 1, label).font = BLACK
        c = ws.cell(r, 2, value)
        c.number_format = fmt
        is_input = not (isinstance(value, str) and value.startswith("="))
        c.font = BLUE if is_input else BLACK
        if is_input:
            c.fill = INPUT_FILL
        ws.cell(r, 3, note).font = Font(name=ARIAL, color="666666")
    ws["A13"] = "Rules"
    ws["A13"].font = BOLD
    rules = [
        "Target value for month n = SIP x ((1+g)^n - 1) / g.",
        "Each month: invest target - value of units held, never below ₹0 (it never sells) and at most the max.",
        "Buy at the close of the first NSE trading day of the month. Price index only, no dividends.",
        "Same-budget version: the SIP is set aside every month; the investment comes out of that cash pile, "
        "the rest earns the liquid rate.",
        "Fixed SIP: the SIP is invested every month, for comparison.",
        "Sources: prices NSE historicalOR/indicesHistory; P/E NSE historicalOR/indicesYield from Nov 2015, "
        "earlier from Kaggle 'NIFTY 50 Historical PE, PB, DIV Yield' (matches NSE exactly where both exist).",
        "Built by research/scripts/render_value_averaging_xlsx.py. Past data, not a forecast or advice.",
    ]
    for k, text in enumerate(rules):
        ws.cell(14 + k, 1, text).font = BLACK
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 60


COLS = [
    # key, header, width, format, group
    ("n", "Month no", 9, "0", None),
    ("date", "Date", 12, "dd-mmm-yyyy", None),
    ("close", "Nifty close", 11, "#,##0.00", None),
    ("pe", "P/E", 7, "0.00", None),
    ("target", "Target value", 14, RUPEE, "va"),
    ("vbefore", "Value before buying", 14, RUPEE, "va"),
    ("gap", "Gap (target - value)", 14, RUPEE, "va"),
    ("invest", "Invest this month", 12, RUPEE, "va"),
    ("action", "Action", 10, "@", "va"),
    ("ubought", "Units bought", 10, "0.0000", "va"),
    ("units", "Total units", 10, "0.0000", "va"),
    ("paid", "Total invested", 13, RUPEE, "va"),
    ("vafter", "Value after buying", 14, RUPEE, "va"),
    ("pile_b", "Cash pile + SIP", 12, RUPEE, "fair"),
    ("finvest", "Invest from pile", 12, RUPEE, "fair"),
    ("pile_a", "Cash pile after", 12, RUPEE, "fair"),
    ("funits", "Total units", 10, "0.0000", "fair"),
    ("fvalue", "Total value (units + cash)", 15, RUPEE, "fair"),
    ("xunits", "Total units", 10, "0.0000", "fixed"),
    ("xpaid", "Total invested", 13, RUPEE, "fixed"),
    ("xvalue", "Value", 14, RUPEE, "fixed"),
    ("year", "Year", 7, "0", None),
    ("cf_va", "Cash flow: VA", 12, RUPEE, "va"),
    ("cf_fair", "Cash flow: same budget", 12, RUPEE, "fair"),
    ("cf_fixed", "Cash flow: fixed", 12, RUPEE, "fixed"),
]
COL = {key: get_column_letter(i + 1) for i, (key, *_rest) in enumerate(COLS)}


def monthly_sheet(wb, ledger):
    ws = wb.create_sheet("Monthly")
    groups = {"va": "Value averaging (rule 20)", "fair": "Same ₹10k budget, cash pile at liquid rate",
              "fixed": "Fixed SIP (comparison)"}
    for i, (_key, _h, _w, _f, g) in enumerate(COLS, start=1):
        if g:
            c = ws.cell(1, i, groups[g])
            c.font = BOLD
            c.fill = PatternFill("solid", fgColor=GROUP_FILL[g])
    for i, (_key, header, width, _fmt, _g) in enumerate(COLS, start=1):
        c = ws.cell(2, i, header)
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[2].height = 32

    S = "Settings!$B$4"
    G = "Settings!$B$5"
    CAP = "Settings!$B$6"
    LM = "Settings!$B$8"
    first, last = 3, 2 + len(ledger)
    for k, row in enumerate(ledger):
        r = first + k
        p = r - 1
        c = {key: f"{COL[key]}{r}" for key in COL}
        pc = {key: f"{COL[key]}{p}" for key in COL}
        f = {
            "n": k + 1,
            "date": date.fromisoformat(row["date"]),
            "close": float(row["nifty_close"]),
            "pe": float(row["pe"]) if row["pe"] else None,
            "target": f"={S}*((1+{G})^{c['n']}-1)/{G}",
            "vbefore": "=0" if k == 0 else f"={pc['units']}*{c['close']}",
            "gap": f"={c['target']}-{c['vbefore']}",
            "invest": f"=MIN({CAP},MAX(0,{c['gap']}))",
            "action": f'=IF({c["invest"]}=0,"Skip",IF({c["invest"]}>={CAP}-0.5,"Max","Partial"))',
            "ubought": f"={c['invest']}/{c['close']}",
            "units": f"={c['ubought']}" if k == 0 else f"={pc['units']}+{c['ubought']}",
            "paid": f"={c['invest']}" if k == 0 else f"={pc['paid']}+{c['invest']}",
            "vafter": f"={c['units']}*{c['close']}",
            "pile_b": f"={S}" if k == 0 else f"={pc['pile_a']}*(1+{LM})+{S}",
            "finvest": (f"=MIN({c['pile_b']},{CAP},MAX(0,{c['target']}))" if k == 0 else
                        f"=MIN({c['pile_b']},{CAP},MAX(0,{c['target']}-{pc['funits']}*{c['close']}))"),
            "pile_a": f"={c['pile_b']}-{c['finvest']}",
            "funits": (f"={c['finvest']}/{c['close']}" if k == 0 else
                       f"={pc['funits']}+{c['finvest']}/{c['close']}"),
            "fvalue": f"={c['funits']}*{c['close']}+{c['pile_a']}",
            "xunits": f"={S}/{c['close']}" if k == 0 else f"={pc['xunits']}+{S}/{c['close']}",
            "xpaid": f"={S}*{c['n']}",
            "xvalue": f"={c['xunits']}*{c['close']}",
            "year": f"=YEAR({c['date']})",
            "cf_va": f"=-{c['invest']}",
            "cf_fair": f"=-{S}",
            "cf_fixed": f"=-{S}",
        }
        for i, (key, _h, _w, fmt, _g) in enumerate(COLS, start=1):
            cell = ws.cell(r, i, f[key])
            cell.number_format = fmt
            cell.font = BLUE if key in ("date", "close", "pe") else BLACK
            cell.border = THIN

    # closing row: what each version is worth on the valuation date
    end = last + 1
    ws.cell(end, 1, "End").font = BOLD
    ws.cell(end, 2, "=Settings!$B$10").number_format = "dd-mmm-yyyy"
    ws.cell(end, 3, "=Settings!$B$11").number_format = "#,##0.00"
    for key, formula in (
        ("vafter", f"={COL['units']}{last}*{COL['close']}{end}"),
        ("fvalue", f"={COL['funits']}{last}*{COL['close']}{end}+{COL['pile_a']}{last}"),
        ("xvalue", f"={COL['xunits']}{last}*{COL['close']}{end}"),
        ("cf_va", f"={COL['vafter']}{end}"),
        ("cf_fair", f"={COL['fvalue']}{end}"),
        ("cf_fixed", f"={COL['xvalue']}{end}"),
    ):
        cell = ws.cell(end, COLS.index(next(x for x in COLS if x[0] == key)) + 1, formula)
        cell.number_format = RUPEE
    for i in range(1, len(COLS) + 1):
        ws.cell(end, i).font = BOLD

    rng = f"A{first}:{get_column_letter(len(COLS))}{last}"
    act = f"${COL['action']}{first}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{act}="Max"'],
                                                   fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{act}="Partial"'],
                                                   fill=PatternFill("solid", fgColor="FFF2CC")))
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(COLS))}{last}"
    return first, last, end


def yearly_sheet(wb, first, last):
    ws = wb.create_sheet("Yearly")
    heads = ["Year", "Months", "Skipped (₹0)", "Max ₹30k", "Partial", "Invested this year",
             "Total invested (end)", "Value (end)", "Target (end)", "Same-budget value (end)",
             "Fixed SIP value (end)", "Nifty (last SIP date)"]
    widths = [8, 8, 10, 10, 9, 14, 15, 15, 15, 16, 16, 12]
    for i, (h, w) in enumerate(zip(heads, widths), start=1):
        c = ws.cell(1, i, h)
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32

    def m(key):
        return f"Monthly!${COL[key]}${first}:${COL[key]}${last}"

    for k, year in enumerate(range(1996, 2027)):
        r = 2 + k
        pos = f"MATCH($A{r},{m('year')},1)"
        cells = [
            (year, "0", BLUE),
            (f"=COUNTIFS({m('year')},$A{r})", "0", BLACK),
            (f'=COUNTIFS({m("year")},$A{r},{m("action")},"Skip")', "0", BLACK),
            (f'=COUNTIFS({m("year")},$A{r},{m("action")},"Max")', "0", BLACK),
            (f'=COUNTIFS({m("year")},$A{r},{m("action")},"Partial")', "0", BLACK),
            (f"=SUMIFS({m('invest')},{m('year')},$A{r})", RUPEE, BLACK),
            (f"=INDEX({m('paid')},{pos})", RUPEE, BLACK),
            (f"=INDEX({m('vafter')},{pos})", RUPEE, BLACK),
            (f"=INDEX({m('target')},{pos})", RUPEE, BLACK),
            (f"=INDEX({m('fvalue')},{pos})", RUPEE, BLACK),
            (f"=INDEX({m('xvalue')},{pos})", RUPEE, BLACK),
            (f"=INDEX({m('close')},{pos})", "#,##0", BLACK),
        ]
        for i, (v, fmt, font) in enumerate(cells, start=1):
            c = ws.cell(r, i, v)
            c.number_format, c.font, c.border = fmt, font, THIN
    tot = 2 + (2026 - 1996) + 1
    ws.cell(tot, 1, "Total").font = BOLD
    for i in range(2, 7):
        col = get_column_letter(i)
        c = ws.cell(tot, i, f"=SUM({col}2:{col}{tot - 1})")
        c.font = BOLD
        c.number_format = RUPEE if i == 6 else "0"
    ws.freeze_panes = "B2"


def summary_sheet(wb, first, end):
    ws = wb.create_sheet("Summary", 0)
    ws["A1"] = "Value averaging vs fixed SIP - Nifty 50, price only"
    ws["A1"].font = Font(name=ARIAL, bold=True, size=14)
    heads = ["Version", "Money put in", "Value on valuation date", "XIRR", "What it means"]
    for i, h in enumerate(heads, start=1):
        c = ws.cell(3, i, h)
        c.font, c.fill = HEAD, HEAD_FILL

    def rng(key):
        return f"Monthly!{COL[key]}{first}:{COL[key]}{end}"

    dates = f"Monthly!{COL['date']}{first}:{COL['date']}{end}"
    rows = [
        ("Value averaging (rule 20)", f"=Monthly!{COL['paid']}{end - 1}", f"=Monthly!{COL['vafter']}{end}",
         f"=XIRR({rng('cf_va')},{dates})", "Counts only the money it invested; idle months are ignored"),
        ("Same budget, idle cash at liquid rate", f"=Monthly!{COL['xpaid']}{end - 1}",
         f"=Monthly!{COL['fvalue']}{end}", f"=XIRR({rng('cf_fair')},{dates})",
         "Fair comparison: the full SIP is set aside every month"),
        ("Fixed SIP", f"=Monthly!{COL['xpaid']}{end - 1}", f"=Monthly!{COL['xvalue']}{end}",
         f"=XIRR({rng('cf_fixed')},{dates})", "₹10k every month, no rule"),
    ]
    for k, (name, paid, value, x, note) in enumerate(rows):
        r = 4 + k
        for i, (v, fmt) in enumerate([(name, "@"), (paid, RUPEE), (value, RUPEE), (x, PCT), (note, "@")],
                                     start=1):
            c = ws.cell(r, i, v)
            c.number_format, c.font, c.border = fmt, BLACK, THIN
    ws["A8"] = "Months with ₹0 invested"
    ws["B8"] = f'=COUNTIF(Monthly!{COL["action"]}{first}:{COL["action"]}{end - 1},"Skip")'
    ws["A9"] = "Months at the ₹30k max"
    ws["B9"] = f'=COUNTIF(Monthly!{COL["action"]}{first}:{COL["action"]}{end - 1},"Max")'
    ws["A10"] = "Months partly invested"
    ws["B10"] = f'=COUNTIF(Monthly!{COL["action"]}{first}:{COL["action"]}{end - 1},"Partial")'
    for r in (8, 9, 10):
        ws.cell(r, 1).font = BLACK
        ws.cell(r, 2).font = BLACK
    ws["A12"] = ("Sheets: Monthly = every month (green = ₹30k max, yellow = partial, white = skipped); "
                 "Yearly = roll-up; Settings = inputs.")
    ws["A12"].font = Font(name=ARIAL, italic=True, color="666666")
    for col, w in zip("ABCDE", [38, 16, 22, 10, 58]):
        ws.column_dimensions[col].width = w


def main():
    with open(LEDGER) as fh:
        ledger = list(csv.DictReader(fh))
    val_date, val_close = last_close()
    wb = Workbook()
    settings_sheet(wb, val_date, val_close)
    first, last, end = monthly_sheet(wb, ledger)
    yearly_sheet(wb, first, last)
    summary_sheet(wb, first, end)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}: {len(ledger)} months, rows {first}-{last}, end row {end}")


if __name__ == "__main__":
    main()
